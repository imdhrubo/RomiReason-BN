"""Source-specific normalization into auditable seed candidates."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any


def normalize_bennumeval_row(config: str, row: dict[str, Any]) -> dict[str, str]:
    """Preserve all task-specific input fields in a single seed text field."""
    question_number = row["Q No."]
    if config == "QNLI":
        text = f"প্রতিজ্ঞা: {row['Premise']}\nঅনুমান: {row['Hypothesis']}"
    else:
        text = str(row["Question"])
        if config == "CQ":
            text += f"\nবিকল্প 1: {row['option1']}\nবিকল্প 2: {row['option2']}"
    return {
        "source_dataset": "BenNumEval",
        "source_config": config,
        "source_item_id": f"{config}-{question_number}",
        "task_type": "math_reasoning",
        "text": text,
        "answer": str(row["Answer"]),
    }


def normalize_bennumeval_rows(
    config: str, rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Normalize a configuration without silently discarding malformed rows."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        try:
            normalized = normalize_bennumeval_row(config, row)
            if not normalized["text"].strip() or not normalized["answer"].strip():
                raise ValueError("BenNumEval row has empty text or answer")
            accepted.append(normalized)
        except (KeyError, ValueError) as error:
            rejected.append(
                {
                    "source_dataset": "BenNumEval",
                    "source_config": config,
                    "source_row_index": str(row_index),
                    "reason": str(error),
                }
            )
    return accepted, rejected


def normalize_bnmmlu_row(row: dict[str, Any]) -> dict[str, str]:
    """Render BnMMLU's options while retaining its original answer label."""
    answer = str(row["correct_answer"]).strip().upper()
    options = list(row["options"])
    labels = ("A", "B", "C", "D")
    if len(options) != len(labels) or answer not in labels:
        raise ValueError("BnMMLU row must have four options and an A-D answer")
    rendered_options = "\n".join(
        f"{label}) {option}" for label, option in zip(labels, options)
    )
    return {
        "source_dataset": "BnMMLU",
        "source_subject": str(row["subject_name"]),
        "source_item_id": str(row["Unique_Serial"]),
        "task_type": "basic_understanding",
        "text": f"{row['question']}\n{rendered_options}",
        "answer": answer,
    }


def normalize_bnmmlu_rows(
    rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Normalize BnMMLU while retaining records that fail structural checks."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        try:
            normalized = normalize_bnmmlu_row(row)
            if not normalized["text"].strip():
                raise ValueError("BnMMLU row has empty text")
            accepted.append(normalized)
        except (KeyError, TypeError, ValueError) as error:
            rejected.append(
                {
                    "source_dataset": "BnMMLU",
                    "source_row_index": str(row_index),
                    "source_item_id": str(row.get("Unique_Serial", "")),
                    "reason": str(error),
                }
            )
    return accepted, rejected


def normalize_ganit_rows(
    rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Keep only Ganit rows marked valid upstream, without repairing answers."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        source_item_id = str(row.get("id", ""))
        try:
            if int(row["valid"]) != 1:
                raise ValueError("Ganit row is not marked valid upstream")
            text = str(row["problem"]).strip()
            answer = str(row["bengali_solution"]).strip()
            if not text or not answer:
                raise ValueError("Ganit row has empty problem or Bengali solution")
            accepted.append(
                {
                    "source_dataset": "Ganit",
                    "source_name": str(row["source_name"]),
                    "source_item_id": source_item_id,
                    "source_difficulty": str(row["difficulty"]),
                    "task_type": "math_reasoning",
                    "text": text,
                    "answer": answer,
                }
            )
        except (KeyError, TypeError, ValueError) as error:
            rejected.append(
                {
                    "source_dataset": "Ganit",
                    "source_row_index": str(row_index),
                    "source_item_id": source_item_id,
                    "reason": str(error),
                }
            )
    return accepted, rejected


def normalize_bluck_csv(
    source_file: Path, rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Normalize BLUCK's lowercase-label MCQ CSVs and retain malformed rows."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        try:
            answer = str(row["answer"]).strip().upper()
            labels = ("A", "B", "C", "D")
            question = str(row["question"]).strip()
            options = [str(row[label.lower()]).strip() for label in labels]
            if answer not in labels or not question or not all(options):
                raise ValueError("BLUCK row has invalid answer or empty text")
            accepted.append(
                {
                    "source_dataset": "BLUCK",
                    "source_file": str(source_file),
                    "source_item_id": f"{source_file}:{row_index}",
                    "source_category": source_file.parent.name,
                    "task_type": "basic_understanding",
                    "text": question + "\n" + "\n".join(
                        f"{label}) {option}" for label, option in zip(labels, options)
                    ),
                    "answer": answer,
                }
            )
        except (KeyError, ValueError) as error:
            rejected.append(
                {
                    "source_dataset": "BLUCK",
                    "source_file": str(source_file),
                    "source_row_index": str(row_index),
                    "reason": str(error),
                }
            )
    return accepted, rejected


def read_bluck_directory(
    root: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Read every benchmark CSV below BLUCK's bn_dataset directory."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for source_file in sorted(root.rglob("*.csv")):
        with source_file.open(encoding="utf-8", newline="") as handle:
            valid, invalid = normalize_bluck_csv(
                source_file.relative_to(root), list(csv.DictReader(handle))
            )
        accepted.extend(valid)
        rejected.extend(invalid)
    return accepted, rejected


def normalize_banglamath_csv(
    rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Normalize BanglaMATH's primary question-answer CSV."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        try:
            text = str(row["Question"]).strip()
            answer = str(row["Answer"]).strip()
            if not text or not answer:
                raise ValueError("BanglaMATH row has empty question or answer")
            accepted.append(
                {
                    "source_dataset": "BanglaMATH",
                    "source_item_id": str(row_index),
                    "source_grade": str(row.get("Grade", "")).strip(),
                    "source_steps": str(row.get("Steps", "")).strip(),
                    "task_type": "math_reasoning",
                    "text": text,
                    "answer": answer,
                }
            )
        except (KeyError, ValueError) as error:
            rejected.append(
                {
                    "source_dataset": "BanglaMATH",
                    "source_row_index": str(row_index),
                    "reason": str(error),
                }
            )
    return accepted, rejected


def read_banglamath_csv(
    input_path: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    with input_path.open(encoding="utf-8-sig", newline="") as handle:
        return normalize_banglamath_csv(list(csv.DictReader(handle)))


def normalize_bmwp_rows(
    rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Normalize BMWP workbook rows, retaining its supplied equation as provenance."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        try:
            text = str(row["Problem"]).strip()
            answer = str(row["Sollution"]).strip()
            if not text or not answer or answer.lower() == "none":
                raise ValueError("BMWP row has empty problem or solution")
            accepted.append(
                {
                    "source_dataset": "BMWP",
                    "source_item_id": str(row_index),
                    "source_equation": str(row.get("Equation", "")).strip(),
                    "source_class": str(row.get("Class", "")).strip(),
                    "task_type": "math_reasoning",
                    "text": text,
                    "answer": answer,
                }
            )
        except (KeyError, ValueError) as error:
            rejected.append(
                {
                    "source_dataset": "BMWP",
                    "source_row_index": str(row_index),
                    "reason": str(error),
                }
            )
    return accepted, rejected


def read_bmwp_workbook(
    input_path: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Read BMWP with openpyxl to preserve its workbook-only source format."""
    from openpyxl import load_workbook

    worksheet = load_workbook(input_path, read_only=True, data_only=True).active
    values = list(worksheet.iter_rows(values_only=True))
    headers = [str(value) if value is not None else "" for value in values[0]]
    rows = [
        {header: value for header, value in zip(headers, values_row)}
        for values_row in values[1:]
    ]
    return normalize_bmwp_rows(rows)


def normalize_benqa_row(source_file: str, row_index: int, row: dict[str, Any]) -> dict[str, str]:
    """Use BEnQA's Bengali question and Bengali MCQ options, never its English fields."""
    answer = str(row["Correct Answer"]).strip().upper()
    labels = ("A", "B", "C", "D")
    if answer not in labels:
        raise ValueError("BEnQA row must have an A-D answer")
    options = [str(row[f"{label} Bn"]).strip() for label in labels]
    if not all(options) or not str(row["Bengali Question"]).strip():
        raise ValueError("BEnQA row has empty Bengali question or option")
    rendered_options = "\n".join(
        f"{label}) {option}" for label, option in zip(labels, options)
    )
    return {
        "source_dataset": "BEnQA",
        "source_file": source_file,
        "source_item_id": f"{source_file}:{row_index}",
        "task_type": "unassigned_pending_expert_task_mapping",
        "text": f"{row['Bengali Question']}\n{rendered_options}",
        "answer": answer,
    }


def normalize_benqa_rows(
    source_file: str, rows: list[dict[str, Any]]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Normalize valid BEnQA rows and retain explicit records for rejected rows."""
    accepted: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        try:
            accepted.append(normalize_benqa_row(source_file, row_index, row))
        except ValueError as error:
            rejected.append(
                {
                    "source_dataset": "BEnQA",
                    "source_file": source_file,
                    "source_item_id": f"{source_file}:{row_index}",
                    "reason": str(error),
                    "raw_answer": str(row.get("Correct Answer", "")),
                }
            )
    return accepted, rejected
