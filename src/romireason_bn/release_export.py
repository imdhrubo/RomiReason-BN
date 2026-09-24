"""Portable Parquet and human-review Excel exports for a frozen dataset."""

from __future__ import annotations

import json
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REVIEW_COLUMNS = (
    "item_id",
    "native_text",
    "answer",
    "task_type",
    "source_dataset",
    "source_item_id",
    "source_revision",
    "source_split",
    "source_config",
    "semantic_cluster_id",
    "semantic_deduplication",
    "source_provenance_json",
)


def _flat_row(row: dict[str, Any]) -> dict[str, str]:
    source = row["source_provenance"]
    semantic = row.get("semantic_dedup_provenance", {})
    return {
        "item_id": row["item_id"],
        "native_text": row["native_text"],
        "answer": str(row["answer"]),
        "task_type": row["task_type"],
        "source_dataset": str(source.get("source_dataset", "")),
        "source_item_id": str(source.get("source_item_id", "")),
        "source_revision": str(source.get("source_revision", "")),
        "source_split": str(source.get("source_split", "")),
        "source_config": str(source.get("source_config", "")),
        "semantic_cluster_id": str(semantic.get("cluster_id", "")),
        "semantic_deduplication": "yes" if semantic else "no",
        "source_provenance_json": json.dumps(source, ensure_ascii=False, sort_keys=True),
    }


def export_release(
    rows: Iterable[dict[str, Any]], parquet_path: Path, workbook_path: Path
) -> dict[str, Any]:
    """Write the same flat, provenance-preserving dataset to Parquet and XLSX."""
    records = [_flat_row(row) for row in rows]
    parquet_path.parent.mkdir(parents=True, exist_ok=True)
    pq.write_table(pa.Table.from_pylist(records), parquet_path, compression="zstd")

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["Dataset review workbook"])
    readme.append(["Rows", len(records)])
    readme.append(["Main sheet", "items"])
    readme.append(["Review use", "Filter by source_dataset or task_type; preserve item_id in feedback."])
    readme.append(["Provenance", "Source fields are flattened; the final column retains full source JSON."])
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 100

    items = workbook.create_sheet("items")
    items.append(list(REVIEW_COLUMNS))
    for record in records:
        items.append([record[column] for column in REVIEW_COLUMNS])
    items.freeze_panes = "A2"
    items.auto_filter.ref = items.dimensions
    for cell in items[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)
    widths = {
        "A": 20, "B": 70, "C": 18, "D": 22, "E": 18, "F": 24,
        "G": 42, "H": 18, "I": 18, "J": 30, "K": 22, "L": 80,
    }
    for column, width in widths.items():
        items.column_dimensions[column].width = width
    for row in items.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    sources = workbook.create_sheet("source_summary")
    sources.append(["source_dataset", "items"])
    for source, count in sorted(Counter(r["source_dataset"] for r in records).items()):
        sources.append([source, count])
    tasks = workbook.create_sheet("task_summary")
    tasks.append(["task_type", "items"])
    for task, count in sorted(Counter(r["task_type"] for r in records).items()):
        tasks.append([task, count])
    for sheet in (sources, tasks):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 14
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")

    workbook.save(workbook_path)
    return {
        "rows": len(records),
        "review_columns": list(REVIEW_COLUMNS),
        "source_counts": dict(Counter(r["source_dataset"] for r in records)),
        "task_counts": dict(Counter(r["task_type"] for r in records)),
    }
